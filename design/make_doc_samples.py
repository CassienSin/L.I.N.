"""Builds the invented sample documents used by voice/verify_docs.py (run once; the files are committed)."""
from pathlib import Path
import sys
out = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).parent / 'doc-samples')
out.mkdir(parents=True, exist_ok=True)
from PIL import Image, ImageDraw, ImageFont
font_path = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
def picture(lines, size=(1400, 500)):
    img = Image.new('RGB', size, 'white'); d = ImageDraw.Draw(img); f = ImageFont.truetype(font_path, 56)
    for i, line in enumerate(lines):
        d.text((60, 60 + i * 90), line, fill='black', font=f)
    return img

import docx
doc = docx.Document(); doc.add_heading('Aurora supply contract', 1)
doc.add_paragraph('The late delivery penalty is 2 percent per week, capped at 10 percent.')
doc.add_paragraph(' '.join(['This long clause exists to test wrapping of very long paragraphs.'] * 30) + ' Wrapmarker zebra appears at the end.')
t = doc.add_table(rows=2, cols=2); t.cell(0, 0).text = 'Item'; t.cell(0, 1).text = 'Price'; t.cell(1, 0).text = 'Solar panel'; t.cell(1, 1).text = '18,500 pesos'
doc.save(out / 'aurora-contract.docx')

from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
c = canvas.Canvas(str(out / 'harbor-report.pdf'), pagesize=A4)
for n, text in enumerate(['Harbor quarterly report', 'Harbor crane uptime reached 97 percent this quarter.', 'Next review is in January.'], 1):
    c.setFont('Helvetica', 14); c.drawString(72, 760, f'Page {n}'); c.drawString(72, 720, text); c.showPage()
c.save()
scan = out / '_scan.png'; picture(['Receipt total 4850 pesos', 'Paid in cash at Aling Nena store'], (1654, 600)).save(scan)
c = canvas.Canvas(str(out / 'scanned-receipt.pdf'), pagesize=A4); c.drawImage(str(scan), 0, 400, width=595, height=216); c.showPage(); c.save(); scan.unlink()
picture(['Sprint goal: ship the voice orb', 'Owner: Lin team']).save(out / 'whiteboard.png')
from reportlab.lib import pdfencrypt
c = canvas.Canvas(str(out / 'locked.pdf'), encrypt=pdfencrypt.StandardEncryption('secret-sample', canPrint=0)); c.drawString(72, 720, 'Locked marker'); c.showPage(); c.save()
(out / 'broken.pdf').write_bytes(b'%PDF-1.4\nthis is not really a pdf\n')
(out / 'old-minutes.doc').write_bytes(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 504)

import openpyxl
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Summary'; ws.append(['Category', 'Amount']); ws.append(['Total', 30500])
ws2 = wb.create_sheet('Travel'); ws2.append(['Flights to Cebu', 12000]); ws2.append(['Hotel', 8000]); wb.save(out / 'budget.xlsx')
from pptx import Presentation
p = Presentation(); s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = 'Lin launch'; s.placeholders[1].text = 'A local assistant'
s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = 'Timeline'; s.placeholders[1].text = 'The launch date is March 14.'
p.save(out / 'pitch.pptx')
from odf.opendocument import OpenDocumentText
from odf.text import P
o = OpenDocumentText(); o.text.addElement(P(text='OpenDocument marker: the kiwi orchard needs water on Fridays.')); o.save(str(out / 'orchard.odt'))
(out / 'memo.rtf').write_text(r'{\rtf1\ansi{\fonttbl\f0\fswiss Helvetica;}\f0\pard Rich text marker: the lantern festival starts at seven.\par}', encoding='ascii')
(out / 'reef.html').write_text('<html><body><h1>Reef notes</h1><p>HTML marker: the coral reef survey is on Tuesday.</p><script>var x=1</script></body></html>', encoding='utf-8')
print(sorted(p.name for p in out.iterdir()))
